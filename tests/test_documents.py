import platform
import shutil
import subprocess

import pymupdf
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from doc2audio.documents import (
    extract_document,
    normalize_ocr_numbers,
    order_ocr_fragments,
    page_indices,
)
from doc2audio.errors import Doc2AudioError

KOREAN = "안녕하세요. 오늘은 한국어 문서를 읽습니다."


def test_ocr_number_spacing_is_fixed_only_with_explicit_unit():
    assert normalize_ocr_numbers("가격은 1 2, 5 0 0 원이며 3 0 분입니다. 목록: 1 2 3") == (
        "가격은 12,500원이며 30분입니다. 목록: 1 2 3"
    )


def test_ocr_split_line_is_reassembled_before_the_next_paragraph():
    assert (
        order_ocr_fragments(
            [
                (0.1, 0.83, 0.02, "예약 시간은 오후 3시"),
                (0.1, 0.78, 0.02, "마지막 문장입니다."),
                (0.7, 0.831, 0.02, "3 0 분입니다."),
            ]
        )
        == "예약 시간은 오후 3시 30분입니다.\n\n마지막 문장입니다."
    )


def make_pdf(path, lines=(KOREAN, "마지막 페이지입니다.")):
    with pymupdf.open() as pdf:
        for line in lines:
            page = pdf.new_page()
            page.insert_text((60, 100), line, fontname="korea", fontsize=18)
        pdf.save(path)


def test_pdf_korean_and_page_selection(tmp_path):
    source = tmp_path / "한글 문서.pdf"
    make_pdf(source)
    document = extract_document(source, ocr="never")
    assert KOREAN in document.text
    assert document.text.endswith("마지막 페이지입니다.")
    assert document.pages == 2
    assert extract_document(source, pages="2", ocr="never").text == "마지막 페이지입니다."


@pytest.mark.parametrize("spec", ["0", "2-1", "3", "1,", "a", "1-999"])
def test_invalid_pdf_page_selection(spec):
    with pytest.raises(Doc2AudioError):
        page_indices(spec, 2)


def test_docx_keeps_body_table_hyperlink_order(tmp_path):
    source = tmp_path / "단락과 표.docx"
    doc = Document()
    first = doc.add_paragraph("시작. ")
    hyperlink = OxmlElement("w:hyperlink")
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "링크 본문"
    run.append(text)
    hyperlink.append(run)
    first._p.append(hyperlink)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "품목"
    table.cell(0, 1).text = "수량"
    table.cell(1, 0).text = "사과"
    table.cell(1, 1).text = "세 개"
    doc.add_paragraph("끝.")
    doc.sections[0].header.paragraphs[0].text = "반복 머리말"
    doc.save(source)
    assert extract_document(source).text == "시작. 링크 본문\n\n품목, 수량.\n\n사과, 세 개.\n\n끝."


def test_docx_tracked_changes_are_not_silently_dropped(tmp_path):
    source = tmp_path / "tracked.docx"
    doc = Document()
    paragraph = doc.add_paragraph("원문")
    change = OxmlElement("w:ins")
    change.set(qn("w:id"), "1")
    paragraph._p.append(change)
    doc.save(source)
    with pytest.raises(Doc2AudioError, match="변경"):
        extract_document(source)


def test_docx_vertical_merged_cell_is_read_once(tmp_path):
    source = tmp_path / "merged.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "공통 항목"
    table.cell(0, 1).text = "첫 값"
    table.cell(1, 1).text = "둘째 값"
    doc.save(source)
    text = extract_document(source).text
    assert text.count("공통 항목") == 1
    assert "첫 값" in text and "둘째 값" in text


@pytest.mark.integration
@pytest.mark.skipif(not shutil.which("textutil"), reason="macOS textutil required")
def test_real_binary_doc_conversion(tmp_path):
    text = tmp_path / "테스트 본문.txt"
    text.write_text(KOREAN, encoding="utf-8")
    source = tmp_path / "테스트 파일.doc"
    subprocess.run(["textutil", "-convert", "doc", "-output", str(source), str(text)], check=True)
    assert source.read_bytes()[:8] == bytes.fromhex("d0cf11e0a1b11ae1")
    assert extract_document(source).text == KOREAN


@pytest.mark.integration
@pytest.mark.skipif(platform.system() != "Darwin", reason="macOS Vision required")
def test_real_scanned_korean_pdf_ocr(tmp_path):
    original = tmp_path / "original.pdf"
    scanned = tmp_path / "스캔.pdf"
    make_pdf(original, [KOREAN])
    with pymupdf.open(original) as pdf, pymupdf.open() as scan:
        image = pdf[0].get_pixmap(matrix=pymupdf.Matrix(2, 2))
        page = scan.new_page()
        page.insert_image(page.rect, stream=image.tobytes("png"))
        scan.save(scanned)
    result = extract_document(scanned)
    assert result.ocr_pages == [1]
    assert "한국어 문서를 읽습니다" in result.text
    with pytest.raises(Doc2AudioError, match="스캔"):
        extract_document(scanned, ocr="never")


def test_encrypted_empty_and_corrupt_documents(tmp_path):
    encrypted = tmp_path / "암호.pdf"
    with pymupdf.open() as pdf:
        pdf.new_page()
        pdf.save(
            encrypted,
            encryption=pymupdf.PDF_ENCRYPT_AES_256,
            owner_pw="test-owner",
            user_pw="test-user",
        )
    with pytest.raises(Doc2AudioError, match="암호화"):
        extract_document(encrypted)
    empty = tmp_path / "empty.txt"
    empty.write_text(" \n\n")
    with pytest.raises(Doc2AudioError, match="본문"):
        extract_document(empty)
    corrupt = tmp_path / "broken.docx"
    corrupt.write_bytes(b"not a document")
    with pytest.raises(Doc2AudioError, match="문서를 읽지 못"):
        extract_document(corrupt)
