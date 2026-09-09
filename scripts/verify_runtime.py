"""Generate small synthetic documents, then narrate them with the REAL local model.

Run: uv run python scripts/verify_runtime.py
This is an opt-in integration check, not part of the fast pytest suite.
No supplied/user documents are read. All generated files stay in output/validation/.
"""

import json
import subprocess
from pathlib import Path

import pymupdf
from docx import Document

from doc2audio.documents import extract_document
from doc2audio.pipeline import convert_document


def main():
    root = Path(__file__).resolve().parents[1]
    output = root / "output" / "validation"
    output.mkdir(parents=True, exist_ok=True)
    pdf_path = output / "한국어 날짜와 금액.pdf"
    pdf_lines = [
        "오늘은 2026년 9월 9일입니다.",
        "가격은 12,500원이며, 예약 시간은 오후 3시 30분입니다.",
        "문서의 마지막 문장까지 빠짐없이 읽습니다.",
    ]
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        for index, line in enumerate(pdf_lines):
            page.insert_text((50, 100 + index * 40), line, fontname="korea", fontsize=13)
        pdf.save(pdf_path)
        pix = page.get_pixmap(matrix=pymupdf.Matrix(2, 2))
        pix.save(output / "pdf-source.png")
        scanned_path = output / "한국어 스캔.pdf"
        with pymupdf.open() as scan:
            scan.new_page().insert_image(page.rect, stream=pix.tobytes("png"))
            scan.save(scanned_path)
    docx_path = output / "한국어 본문과 표.docx"
    docx = Document()
    docx.add_paragraph("안녕하세요. 워드 문서의 본문과 표를 차례로 읽습니다.")
    table = docx.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "사과"
    table.cell(0, 1).text = "세 개"
    docx.add_paragraph("마지막으로, 오늘도 편안한 하루 보내세요.")
    docx.save(docx_path)
    doc_path = output / "한국어 구형 문서.doc"
    text_path = output / "legacy-source.txt"
    text_path.write_text(
        "이 파일은 구형 워드 문서입니다. 한글과 띄어쓰기를 그대로 읽습니다.", encoding="utf-8"
    )
    subprocess.run(
        ["textutil", "-convert", "doc", "-output", str(doc_path), str(text_path)],
        check=True,
    )
    results = []
    for source in (pdf_path, docx_path, doc_path, scanned_path):
        extracted = extract_document(source)
        if source in (pdf_path, scanned_path):
            # Check OCR against the original fixture, not just its own transcript.
            assert "".join(extracted.text.split()) == "".join("".join(pdf_lines).split())
        destination = source.with_suffix(".mp3")
        result = convert_document(source, destination, offline=True, overwrite=True)
        probe = subprocess.run(
            [
                "ffprobe",
                "-v",
                "error",
                "-show_entries",
                "format=duration,size:stream=codec_name,sample_rate,channels",
                "-of",
                "json",
                str(destination),
            ],
            capture_output=True,
            text=True,
            check=True,
        )
        result["input"] = str(source)
        result["extracted_text"] = extracted.text
        result["ffprobe"] = json.loads(probe.stdout)
        assert result["ffprobe"]["streams"][0]["codec_name"] == "mp3"
        assert float(result["ffprobe"]["format"]["duration"]) > 1
        results.append(result)
        (output / "runtime-results.json").write_text(
            json.dumps(results, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
    print(f"REAL_RUNTIME_PASS: {len(results)} document formats; {output / 'runtime-results.json'}")


if __name__ == "__main__":
    main()
